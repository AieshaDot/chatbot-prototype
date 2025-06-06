import os
from pathlib import Path
import pandas as pd
from docx import Document as DocxDocument
import fitz  # PyMuPDF

from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from langchain_community.vectorstores import Chroma
from langchain_community.embeddings import OpenAIEmbeddings
from langchain_community.document_loaders import UnstructuredFileLoader

PERSIST_DIR = "chroma_db"

# ✅ PDF Extractor
def extract_text_from_pdf(file_path):
    doc = fitz.open(file_path)
    text = "\n".join([page.get_text() for page in doc])
    return text.strip()

# ✅ DOCX Extractor
def extract_text_from_docx(file_path):
    doc = DocxDocument(file_path)
    return "\n".join([p.text for p in doc.paragraphs]).strip()

# ✅ XLSX Extractor
def extract_text_from_xlsx(file_path):
    df = pd.read_excel(file_path)
    return df.to_csv(index=False)

def extract_text_from_txt(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()

# ✅ Embed + Store in Vector DB
def load_and_index_file(file_path):
    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"{file_path} not found")

    print(f"📄 Loading: {file_path.name}")

    # Use Unstructured loader (for PDF/DOCX/TXT/HTML/etc.)
    loader = UnstructuredFileLoader(str(file_path))
    raw_docs = loader.load()

    # Split into chunks
    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=200)
    split_docs = splitter.split_documents(raw_docs)

    # Add source metadata
    docs_with_meta = [
        Document(page_content=doc.page_content, metadata={"source": file_path.name})
        for doc in split_docs
    ]

    # Save to Chroma
    vectordb = Chroma.from_documents(
        docs_with_meta,
        embedding=OpenAIEmbeddings(),
        persist_directory=PERSIST_DIR
    )

    vectordb.persist()
    print(f"✅ Indexed and saved {len(docs_with_meta)} chunks from {file_path.name}")
    return len(docs_with_meta)