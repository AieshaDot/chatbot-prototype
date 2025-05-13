from langchain_openai import OpenAIEmbeddings
from langchain_chroma import Chroma
from langchain_community.document_loaders import TextLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema import Document
from typing import List
import os

# 👇 Your persisted directory and custom collection
persist_directory = "db"
collection_name = "itsmf_docs"

# 👇 Embedding function
embedding = OpenAIEmbeddings()

# 👇 Vector DB instance (with collection name preserved)
vectordb = Chroma(
    collection_name=collection_name,
    persist_directory=persist_directory,
    embedding_function=embedding
)

# ✅ Exported retriever for /ask route
retriever = vectordb.as_retriever()

# 👇 Add a document to the collection and persist it
from langchain_community.document_loaders import (
    TextLoader, CSVLoader, PyPDFLoader, UnstructuredExcelLoader
)
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema import Document
from typing import List
from docx import Document as DocxDocument
from bs4 import BeautifulSoup
import os

def load_docx(file_path):
    from docx import Document as DocxDocument
    doc = DocxDocument(file_path)
    text = "\n".join([p.text for p in doc.paragraphs])
    return [Document(page_content=text, metadata={"source": os.path.basename(file_path)})]

def load_html(file_path):
    from bs4 import BeautifulSoup
    with open(file_path, "r", encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    text = soup.get_text(separator="\n")
    return [Document(page_content=text, metadata={"source": os.path.basename(file_path)})]

from langchain.schema import Document
import pandas as pd
import os

def load_xlsx(file_path):
    df = pd.read_excel(file_path)
    text = df.to_csv(index=False)

    return [Document(
        page_content=text,
        metadata={"source": os.path.basename(file_path)}
    )]

def add_documents(file_path: str) -> int:
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        loader = PyPDFLoader(file_path)
        docs = loader.load()
    elif ext == ".docx":
        docs = load_docx(file_path)
    elif ext == ".html":
        docs = load_html(file_path)
    elif ext == ".txt":
        loader = TextLoader(file_path)
        docs = loader.load()
    elif ext == ".csv":
        loader = CSVLoader(file_path)
        docs = loader.load()
    elif ext == ".xlsx":
        docs = load_xlsx(file_path)
    else:
        raise ValueError(f"❌ Unsupported file type: {ext}")

    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=100)
    split_docs = splitter.split_documents(docs)

    vectordb.add_documents(split_docs)
    if hasattr(vectordb, "persist"):
        vectordb.persist()

    return len(split_docs)

# 👇 Run a similarity search on the collection
def query_similar_documents(query: str, k: int = 4) -> List[Document]:
    return vectordb.similarity_search(query, k=k)

vectorstore = vectordb 